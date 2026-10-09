
import asyncio
import base64
import csv
import html
import io
import logging
import os
import re
import warnings
from pathlib import Path

import cv2
import numpy as np
import pytesseract
import pymupdf

from PIL import Image, ImageOps, UnidentifiedImageError
from fastapi import FastAPI, File, UploadFile
from fastapi.responses import HTMLResponse, PlainTextResponse, Response
from fastapi.concurrency import run_in_threadpool

from openpyxl import load_workbook
from docx import Document

try:
    import xlrd
except ImportError:
    xlrd = None

try:
    from document_scanner import scan_document
except ImportError:
    scan_document = None

try:
    from preprocessing import preprocess_image
except ImportError:
    preprocess_image = None

try:
    from smart_extraction import extract_information
except ImportError:
    extract_information = None


# =========================================================
# CONFIGURATION
# =========================================================

APP_TITLE = "AI Document Scanner"

MAX_UPLOAD_BYTES = 10 * 1024 * 1024
MAX_IMAGE_PIXELS = 20_000_000
MAX_PDF_PAGES = 10
MAX_CONCURRENT_SCANS = 2
OCR_TIMEOUT_SECONDS = 15
MAX_EXTRACTED_CHARS = 2_000_000
MAX_SPREADSHEET_CELLS = 100_000

ALLOWED_EXTENSIONS = {
    ".pdf",
    ".png", ".jpg", ".jpeg", ".tif", ".tiff", ".bmp",
    ".xlsx", ".xls", ".csv",
    ".docx",
    ".txt",
}

IMAGE_EXTENSIONS = {
    ".png", ".jpg", ".jpeg", ".tif", ".tiff", ".bmp"
}

logging.basicConfig(
    level=os.getenv("LOG_LEVEL", "INFO").upper(),
    format="%(asctime)s %(levelname)s %(name)s: %(message)s",
)

logger = logging.getLogger("ai_document_scanner")

tesseract_cmd = os.getenv("TESSERACT_CMD")
if tesseract_cmd:
    pytesseract.pytesseract.tesseract_cmd = tesseract_cmd
else:
    default_tesseract = Path(
        r"C:\Program Files\Tesseract-OCR\tesseract.exe"
    )
    if default_tesseract.exists():
        pytesseract.pytesseract.tesseract_cmd = str(default_tesseract)

app = FastAPI(
    title=APP_TITLE,
    docs_url=None,
    redoc_url=None,
    openapi_url=None,
)

scan_semaphore = asyncio.Semaphore(MAX_CONCURRENT_SCANS)


# =========================================================
# SHARED HELPERS
# =========================================================

def extract_smart_information(text):
    if extract_information is None:
        return {}

    try:
        result = extract_information(text)
        return result if isinstance(result, dict) else {}
    except Exception:
        logger.exception("Smart information extraction failed")
        return {}


def make_text_result(text, preview="", processed_preview=""):
    text = (text or "")[:MAX_EXTRACTED_CHARS]

    return {
        "text": text or "No text could be extracted from this document.",
        "information": extract_smart_information(text),
        "original_preview": preview,
        "processed_preview": processed_preview,
    }


def image_to_jpeg_base64(image, max_width=1200):
    if image is None:
        return ""

    try:
        if len(image.shape) == 2:
            image = cv2.cvtColor(image, cv2.COLOR_GRAY2BGR)

        height, width = image.shape[:2]

        if width > max_width:
            new_height = max(1, int(height * max_width / width))
            image = cv2.resize(
                image,
                (max_width, new_height),
                interpolation=cv2.INTER_AREA,
            )

        success, encoded = cv2.imencode(
            ".jpg",
            image,
            [int(cv2.IMWRITE_JPEG_QUALITY), 80],
        )

        if not success:
            return ""

        return base64.b64encode(encoded.tobytes()).decode("ascii")

    except Exception:
        logger.exception("Could not generate preview")
        return ""


def check_file_size(data):
    if not data:
        raise ValueError("The uploaded file is empty.")

    if len(data) > MAX_UPLOAD_BYTES:
        raise ValueError("File is too large. Maximum size is 10 MB.")


def get_image_bytes(image):
    """Convert an RGB Pillow image to an OpenCV BGR array."""
    rgb = np.asarray(image.convert("RGB"))
    return cv2.cvtColor(rgb, cv2.COLOR_RGB2BGR)


def run_ocr(image):
    try:
        return pytesseract.image_to_string(
            image,
            config="--oem 3 --psm 6",
            timeout=OCR_TIMEOUT_SECONDS,
        ).strip()

    except pytesseract.TesseractNotFoundError:
        logger.exception("Tesseract executable was not found")
        raise RuntimeError(
            "Tesseract OCR is not installed or configured. "
            "Install Tesseract or set the TESSERACT_CMD environment variable."
        )

    except RuntimeError:
        logger.warning("OCR operation timed out")
        return ""

    except Exception:
        logger.exception("OCR failed")
        return ""


def process_cv_image(original):
    scanned = original.copy()

    if scan_document is not None:
        try:
            candidate = scan_document(original.copy())
            if isinstance(candidate, np.ndarray) and candidate.size:
                scanned = candidate
        except Exception:
            logger.exception("Document scanning failed; using original image")

    processed = scanned.copy()

    if preprocess_image is not None:
        try:
            candidate = preprocess_image(scanned.copy())
            if isinstance(candidate, np.ndarray) and candidate.size:
                processed = candidate
        except Exception:
            logger.exception("Image preprocessing failed")

    direct_text = run_ocr(scanned)
    processed_text = run_ocr(processed)

    text = max(
        [direct_text, processed_text],
        key=lambda item: len(re.sub(r"\s+", "", item)),
    )

    return {
        "text": text,
        "original_preview": image_to_jpeg_base64(original),
        "processed_preview": image_to_jpeg_base64(processed),
    }


# =========================================================
# IMAGE PROCESSING
# =========================================================

def process_image(data, filename):
    check_file_size(data)

    extension = Path(filename).suffix.lower()
    if extension not in IMAGE_EXTENSIONS:
        raise ValueError("Unsupported image format.")

    try:
        with warnings.catch_warnings():
            warnings.simplefilter("error", Image.DecompressionBombWarning)

            with Image.open(io.BytesIO(data)) as opened:
                if opened.format not in {"PNG", "JPEG", "TIFF", "BMP"}:
                    raise ValueError("The uploaded image format is unsupported.")

                width, height = opened.size
                if width <= 0 or height <= 0:
                    raise ValueError("The image dimensions are invalid.")

                if width * height > MAX_IMAGE_PIXELS:
                    raise ValueError(
                        "Image dimensions are too large. Use a smaller image."
                    )

                opened.verify()

            with Image.open(io.BytesIO(data)) as opened:
                corrected = ImageOps.exif_transpose(opened)
                corrected.load()

                if corrected.width * corrected.height > MAX_IMAGE_PIXELS:
                    raise ValueError(
                        "Image dimensions are too large. Use a smaller image."
                    )

                original = get_image_bytes(corrected)

    except ValueError:
        raise
    except (
        UnidentifiedImageError,
        OSError,
        Image.DecompressionBombError,
        Image.DecompressionBombWarning,
    ):
        raise ValueError("The uploaded image is invalid or too large.")
    except Exception:
        logger.exception("Image decoding failed")
        raise ValueError("Could not open this image.")

    result = process_cv_image(original)

    return make_text_result(
        result["text"],
        result["original_preview"],
        result["processed_preview"],
    )


# =========================================================
# PDF PROCESSING
# =========================================================

def process_pdf(data):
    check_file_size(data)

    try:
        with pymupdf.open(stream=data, filetype="pdf") as pdf:
            if pdf.is_encrypted:
                raise ValueError(
                    "Password-protected PDFs are not supported."
                )

            if pdf.page_count < 1:
                raise ValueError("The PDF has no pages.")

            if pdf.page_count > MAX_PDF_PAGES:
                raise ValueError(
                    f"PDFs can contain at most {MAX_PDF_PAGES} pages."
                )

            pages_text = []
            original_preview = ""
            processed_preview = ""

            for index in range(pdf.page_count):
                page = pdf.load_page(index)

                if page.rect.width <= 0 or page.rect.height <= 0:
                    raise ValueError("A PDF page has invalid dimensions.")

                if page.rect.width * page.rect.height > 2_000_000:
                    raise ValueError(
                        "A PDF page is too large to process safely."
                    )

                # Use embedded text when it is available.
                text = page.get_text("text", sort=True).strip()

                # Render first page for the document preview.
                # Other pages are rendered only when OCR is needed.
                if index == 0 or len(text) < 20:
                    pix = page.get_pixmap(
                        matrix=pymupdf.Matrix(1.5, 1.5),
                        alpha=False,
                    )

                    if pix.width * pix.height > MAX_IMAGE_PIXELS:
                        raise ValueError(
                            "A rendered PDF page is too large."
                        )

                    rgb = np.frombuffer(
                        pix.samples, dtype=np.uint8
                    ).reshape(pix.height, pix.width, 3)

                    image = cv2.cvtColor(rgb, cv2.COLOR_RGB2BGR)

                    if index == 0:
                        original_preview = image_to_jpeg_base64(image)

                    # OCR only when embedded text is absent.
                    if len(text) < 20:
                        processed = image.copy()

                        if scan_document is not None:
                            try:
                                candidate = scan_document(image.copy())
                                if (
                                    isinstance(candidate, np.ndarray)
                                    and candidate.size
                                ):
                                    processed = candidate
                            except Exception:
                                logger.exception(
                                    "PDF page scan failed; using rendered image"
                                )

                        if preprocess_image is not None:
                            try:
                                candidate = preprocess_image(processed.copy())
                                if (
                                    isinstance(candidate, np.ndarray)
                                    and candidate.size
                                ):
                                    processed = candidate
                            except Exception:
                                logger.exception(
                                    "PDF page preprocessing failed"
                                )

                        text = run_ocr(processed)

                        if index == 0:
                            processed_preview = image_to_jpeg_base64(
                                processed
                            )

                pages_text.append(
                    f"Page {index + 1}\n{text}" if text else
                    f"Page {index + 1}\nNo text detected."
                )

                if sum(len(item) for item in pages_text) > MAX_EXTRACTED_CHARS:
                    raise ValueError(
                        "Extracted PDF text exceeds the allowed limit."
                    )

            return make_text_result(
                "\n\n".join(pages_text),
                original_preview,
                processed_preview or original_preview,
            )

    except ValueError:
        raise
    except RuntimeError:
        raise
    except Exception:
        logger.exception("PDF processing failed")
        raise ValueError(
            "Could not read this PDF. Check that it is a valid, unencrypted PDF."
        )


# =========================================================
# EXCEL AND CSV PROCESSING
# =========================================================

def process_excel(data, extension):
    check_file_size(data)
    lines = []

    try:
        if extension == ".csv":
            # UTF-8 with BOM support.
            text = data.decode("utf-8-sig")
            reader = csv.reader(io.StringIO(text))

            for row_number, row in enumerate(reader, start=1):
                if row_number > MAX_SPREADSHEET_CELLS:
                    raise ValueError("CSV contains too many rows.")

                if row:
                    lines.append(
                        " | ".join(str(cell) for cell in row)
                    )

        elif extension == ".xlsx":
            workbook = load_workbook(
                io.BytesIO(data),
                read_only=True,
                data_only=True,
            )

            try:
                cell_count = 0

                for sheet in workbook.worksheets:
                    lines.append(f"Sheet: {sheet.title}")

                    for row in sheet.iter_rows(values_only=True):
                        cell_count += len(row)

                        if cell_count > MAX_SPREADSHEET_CELLS:
                            raise ValueError(
                                "Excel workbook contains too many cells."
                            )

                        values = [
                            "" if value is None else str(value)
                            for value in row
                        ]

                        if any(values):
                            lines.append(" | ".join(values))
            finally:
                workbook.close()

        elif extension == ".xls":
            if xlrd is None:
                raise RuntimeError(
                    "Legacy Excel support is unavailable. "
                    "Install the xlrd package."
                )

            workbook = xlrd.open_workbook(file_contents=data)
            cell_count = 0

            for sheet in workbook.sheets():
                lines.append(f"Sheet: {sheet.name}")

                for row_number in range(sheet.nrows):
                    row = sheet.row_values(row_number)
                    cell_count += len(row)

                    if cell_count > MAX_SPREADSHEET_CELLS:
                        raise ValueError(
                            "Excel workbook contains too many cells."
                        )

                    if any(value != "" for value in row):
                        lines.append(
                            " | ".join(str(value) for value in row)
                        )

        text = "\n".join(lines)

        if len(text) > MAX_EXTRACTED_CHARS:
            raise ValueError("Extracted spreadsheet text is too large.")

        return make_text_result(text)

    except (ValueError, RuntimeError):
        raise
    except Exception:
        logger.exception("Spreadsheet processing failed")
        raise ValueError(
            "Could not read this spreadsheet. Check the file format."
        )


# =========================================================
# WORD PROCESSING
# =========================================================

def process_word(data):
    check_file_size(data)

    try:
        document = Document(io.BytesIO(data))
        lines = []

        for paragraph in document.paragraphs:
            text = paragraph.text.strip()
            if text:
                lines.append(text)

        for table_number, table in enumerate(document.tables, start=1):
            lines.append(f"Table {table_number}:")

            for row in table.rows:
                lines.append(
                    " | ".join(cell.text.strip() for cell in row.cells)
                )

                if sum(len(item) for item in lines) > MAX_EXTRACTED_CHARS:
                    raise ValueError(
                        "Extracted Word document text is too large."
                    )

        text = "\n".join(lines)

        if len(text) > MAX_EXTRACTED_CHARS:
            raise ValueError("Extracted Word document text is too large.")

        return make_text_result(text)

    except ValueError:
        raise
    except Exception:
        logger.exception("Word document processing failed")
        raise ValueError(
            "Could not read this Word document. "
            "Please upload a valid DOCX file."
        )


# =========================================================
# TEXT FILE PROCESSING
# =========================================================

def process_text_file(data):
    check_file_size(data)

    try:
        text = data.decode("utf-8-sig")
    except UnicodeDecodeError:
        try:
            text = data.decode("cp1252")
        except UnicodeDecodeError:
            raise ValueError(
                "Could not decode this text file. Save it as UTF-8."
            )

    if len(text) > MAX_EXTRACTED_CHARS:
        raise ValueError("The text file contains too much text.")

    return make_text_result(text)


# =========================================================
# FILE ROUTER
# =========================================================

def process_uploaded_document(data, filename):
    extension = Path(filename).suffix.lower()

    if extension not in ALLOWED_EXTENSIONS:
        raise ValueError(
            "Unsupported file type. Upload PDF, images, Excel, CSV, "
            "DOCX or TXT files."
        )

    if extension == ".pdf":
        return process_pdf(data)

    if extension in IMAGE_EXTENSIONS:
        return process_image(data, filename)

    if extension in {".xlsx", ".xls", ".csv"}:
        return process_excel(data, extension)

    if extension == ".docx":
        return process_word(data)

    if extension == ".txt":
        return process_text_file(data)

    raise ValueError("This file type is not supported.")


# =========================================================
# WEBSITE HTML
# =========================================================

def make_page(
    result="",
    error="",
    information=None,
    preview="",
    processed_preview="",
):
    error_html = ""
    preview_html = ""
    processed_html = ""
    result_html = ""
    information_html = ""

    if error:
        error_html = f"""
        <div class="message error">
            <strong>Unable to process document</strong>
            <p>{html.escape(error)}</p>
        </div>
        """

    if preview:
        preview_html = f"""
        <div class="preview-card">
            <h3>Original Document</h3>
            <img class="document-preview"
                 src="data:image/jpeg;base64,{preview}"
                 alt="Original document preview">
        </div>
        """

    if processed_preview:
        processed_html = f"""
        <div class="preview-card">
            <h3>Processed Document</h3>
            <img class="document-preview"
                 src="data:image/jpeg;base64,{processed_preview}"
                 alt="Processed document preview">
        </div>
        """

    if preview or processed_preview:
        preview_html = f"""
        <section class="panel">
            <h2>Document Preview</h2>
            <div class="preview-grid">
                {preview_html}
                {processed_html}
            </div>
        </section>
        """

    if result:
        result_html = f"""
        <section class="panel result-panel">
            <div class="section-heading">
                <div>
                    <h2>Extracted Text</h2>
                    <p>Review, copy or download your extracted text.</p>
                </div>
                <span class="status">Processing complete</span>
            </div>

            <textarea id="extractedText"
                      spellcheck="false"
                      aria-label="Extracted document text">{html.escape(result)}</textarea>

            <div class="actions">
                <button type="button" class="button secondary"
                        onclick="copyText()">Copy Text</button>

                <button type="button" class="button primary"
                        onclick="downloadText()">Download Text</button>
            </div>
        </section>
        """

    if information is not None:
        fields = [
            ("Emails", ("emails", "email")),
            ("Phone Numbers", ("phone_numbers", "phones", "phone")),
            ("Dates", ("dates", "date")),
        ]

        cards = []

        for label, keys in fields:
            value = next(
                (information[key] for key in keys if key in information),
                None,
            )

            if value is None or value == "" or value == []:
                display = "None found"
            elif isinstance(value, (list, tuple, set)):
                display = ", ".join(str(item) for item in value)
            else:
                display = str(value)

            cards.append(f"""
                <div class="info-card">
                    <span>{html.escape(label)}</span>
                    <p>{html.escape(display)}</p>
                </div>
            """)

        information_html = f"""
        <section class="panel">
            <div class="section-heading">
                <div>
                    <h2>Smart Information Extraction</h2>
                    <p>Important details identified in your document.</p>
                </div>
            </div>
            <div class="info-grid">
                {''.join(cards)}
            </div>
        </section>
        """

    return f"""<!DOCTYPE html>
<html lang="en">
<head>
<meta charset="UTF-8">
<meta name="viewport" content="width=device-width, initial-scale=1">
<meta name="referrer" content="no-referrer">
<title>{APP_TITLE}</title>
<style>
:root {{
 --primary:#4f46e5; --primary-dark:#4338ca;
 --background:#f4f6fb; --surface:#fff;
 --text:#1f2937; --muted:#6b7280;
 --border:#e5e7eb; --success:#047857; --danger:#b91c1c;
}}
* {{box-sizing:border-box}}
body {{
 margin:0;font-family:"Segoe UI",Arial,sans-serif;
 background:var(--background);color:var(--text);line-height:1.6;
}}
header {{
 background:linear-gradient(135deg,#4338ca,#6366f1);
 color:white;padding:27px 20px;text-align:center;
}}
header h1 {{margin:0;font-size:29px}}
header p {{margin:7px 0 0;opacity:.92}}
main {{width:min(940px,calc(100% - 32px));margin:30px auto 48px}}
.intro {{text-align:center;margin-bottom:24px}}
.intro h2 {{margin-bottom:6px;font-size:25px}}
.intro p,.section-heading p {{color:var(--muted);margin-top:4px}}
.panel {{
 background:var(--surface);border:1px solid var(--border);
 border-radius:14px;padding:25px;margin-bottom:22px;
 box-shadow:0 5px 20px rgba(15,23,42,.045);
}}
.upload-area {{
 border:2px dashed #c7d2fe;background:#fafaff;
 border-radius:12px;padding:35px 20px;text-align:center;
}}
.upload-icon {{font-size:42px;margin-bottom:5px}}
.upload-area h3 {{margin:5px 0}}
.upload-area p {{color:var(--muted);font-size:14px}}
input[type=file] {{
 display:block;width:100%;max-width:500px;margin:20px auto;
 padding:11px;border:1px solid var(--border);border-radius:8px;background:white;
}}
.button {{
 display:inline-block;border:0;border-radius:8px;padding:11px 19px;
 font-size:15px;font-weight:600;cursor:pointer;
}}
.primary {{background:var(--primary);color:white}}
.primary:hover {{background:var(--primary-dark)}}
.secondary {{background:#eef2ff;color:#3730a3}}
.secondary:hover {{background:#e0e7ff}}
.button:disabled {{opacity:.6;cursor:not-allowed}}
.message {{border-radius:9px;padding:14px 17px;margin-bottom:20px}}
.error {{color:var(--danger);background:#fef2f2;border:1px solid #fecaca}}
.error p {{margin-bottom:0}}
.section-heading {{
 display:flex;justify-content:space-between;align-items:center;
 gap:14px;margin-bottom:18px;
}}
.section-heading h2 {{margin:0;font-size:21px}}
.status {{
 color:var(--success);background:#ecfdf5;border-radius:20px;
 padding:5px 11px;font-size:12px;white-space:nowrap;
}}
textarea {{
 width:100%;min-height:230px;resize:vertical;padding:15px;
 border:1px solid #d1d5db;border-radius:9px;
 font:14px/1.65 Consolas,monospace;color:#111827;background:#fcfcfd;
}}
.actions {{display:flex;align-items:center;gap:10px;flex-wrap:wrap;margin-top:15px}}
.info-grid,.preview-grid {{
 display:grid;grid-template-columns:repeat(3,minmax(0,1fr));gap:14px;
}}
.preview-grid {{grid-template-columns:repeat(2,minmax(0,1fr))}}
.info-card,.preview-card {{
 background:#f9fafb;border:1px solid var(--border);
 border-radius:10px;padding:16px;min-width:0;overflow-wrap:anywhere;
}}
.info-card span {{font-size:13px;color:var(--muted)}}
.info-card p {{margin-bottom:0;font-weight:600}}
.preview-card h3 {{margin-top:0;font-size:16px}}
.document-preview {{
 display:block;width:100%;max-height:420px;object-fit:contain;
 background:#f9fafb;border-radius:6px;
}}
.preview-note,footer {{font-size:13px;color:var(--muted);text-align:center}}
footer {{padding:18px}}
@media(max-width:620px) {{
 header h1 {{font-size:24px}}
 main {{margin-top:20px}}
 .panel {{padding:17px}}
 .info-grid,.preview-grid {{grid-template-columns:1fr}}
 .section-heading {{align-items:flex-start;flex-direction:column}}
}}
</style>
</head>
<body>
<header>
 <h1>AI Document Scanner</h1>
 <p>Scan, straighten and extract text from documents</p>
</header>
<main>
 <div class="intro">
  <h2>Smart Document Scanning</h2>
  <p>Upload a document to extract readable text and important information.</p>
 </div>
 {error_html}
 <section class="panel">
  <div class="section-heading">
   <div>
    <h2>Upload Your Document</h2>
    <p>PDF, PNG, JPG, JPEG, TIFF, BMP, XLSX, XLS, CSV, DOCX, TXT</p>
   </div>
  </div>
  <form id="uploadForm" action="/scan" method="post"
        enctype="multipart/form-data">
   <div class="upload-area">
    <div class="upload-icon" aria-hidden="true">📄</div>
    <h3>Select a document</h3>
    <p>Maximum file size: 10 MB</p>
    <input type="file" id="document" name="file"
           accept=".pdf,.png,.jpg,.jpeg,.tif,.tiff,.bmp,.xlsx,.xls,.csv,.docx,.txt"
           required>
    <button id="scanButton" type="submit" class="button primary">
     Scan Document
    </button>
    <p id="uploadStatus" role="status"></p>
   </div>
  </form>
 </section>
 {preview_html}
 {result_html}
 {information_html}
 <p class="preview-note">
  Uploaded files are processed in memory and are not intentionally saved
  as shared documents or scan-history files.
 </p>
</main>
<footer>AI Document Scanner · OCR and document text extraction</footer>
<script>
const form=document.getElementById("uploadForm");
const input=document.getElementById("document");
const button=document.getElementById("scanButton");
const statusText=document.getElementById("uploadStatus");
const maxSize=10*1024*1024;
const allowed=[
".pdf",".png",".jpg",".jpeg",".tif",".tiff",".bmp",
".xlsx",".xls",".csv",".docx",".txt"
];

form.addEventListener("submit",function(event){{
 const file=input.files[0];
 if(!file){{
  event.preventDefault();
  statusText.textContent="Please select a document.";
  return;
 }}
 const dot=file.name.lastIndexOf(".");
 const extension=dot>=0?file.name.substring(dot).toLowerCase():"";
 if(!allowed.includes(extension)){{
  event.preventDefault();
  statusText.textContent="This file type is not supported.";
  return;
 }}
 if(file.size>maxSize){{
  event.preventDefault();
  statusText.textContent="File is too large. Maximum size is 10 MB.";
  return;
 }}
 button.disabled=true;
 button.textContent="Processing...";
 statusText.textContent="Please wait while your document is processed.";
}});

async function copyText(){{
 const field=document.getElementById("extractedText");
 if(!field)return;
 try{{
  await navigator.clipboard.writeText(field.value);
  alert("Text copied successfully.");
 }}catch(error){{
  field.focus();field.select();
  const copied=document.execCommand("copy");
  alert(copied?"Text copied successfully.":"Copy was blocked by your browser.");
 }}
}}

function downloadText(){{
 const field=document.getElementById("extractedText");
 if(!field)return;
 const blob=new Blob([field.value],{{type:"text/plain;charset=utf-8"}});
 const url=URL.createObjectURL(blob);
 const link=document.createElement("a");
 link.href=url;link.download="extracted_text.txt";
 document.body.appendChild(link);link.click();link.remove();
 URL.revokeObjectURL(url);
}}
</script>
</body>
</html>"""


# =========================================================
# RESPONSE SECURITY HEADERS
# =========================================================

def secure_html(content, status_code=200):
    response = HTMLResponse(content, status_code=status_code)

    response.headers["Cache-Control"] = "no-store, max-age=0"
    response.headers["Pragma"] = "no-cache"
    response.headers["X-Content-Type-Options"] = "nosniff"
    response.headers["Referrer-Policy"] = "no-referrer"
    response.headers["Content-Security-Policy"] = (
        "default-src 'self' data:; "
        "img-src 'self' data:; "
        "style-src 'unsafe-inline'; "
        "script-src 'unsafe-inline'; "
        "form-action 'self'; "
        "base-uri 'none'; "
        "frame-ancestors 'none'"
    )

    return response


# =========================================================
# ROUTES
# =========================================================

@app.get("/", response_class=HTMLResponse)
async def home():
    return secure_html(make_page())


@app.post("/scan", response_class=HTMLResponse)
async def scan_document_route(file: UploadFile = File(...)):
    try:
        filename = file.filename or ""
        extension = Path(filename).suffix.lower()

        if extension not in ALLOWED_EXTENSIONS:
            return secure_html(
                make_page(
                    error="Unsupported file type. Upload PDF, images, "
                          "Excel, CSV, DOCX or TXT files."
                ),
                400,
            )

        # Read at most the allowed amount plus one byte.
        data = await file.read(MAX_UPLOAD_BYTES + 1)

        if len(data) > MAX_UPLOAD_BYTES:
            return secure_html(
                make_page(error="File is too large. Maximum size is 10 MB."),
                413,
            )

        async with scan_semaphore:
            try:
                result = await run_in_threadpool(
                    process_uploaded_document,
                    data,
                    filename,
                )

            except ValueError as exc:
                return secure_html(make_page(error=str(exc)), 400)

            except RuntimeError as exc:
                logger.error("Document processor configuration error: %s", exc)
                return secure_html(
                    make_page(error=str(exc)),
                    500,
                )

            except Exception:
                logger.exception("Unexpected document processing error")
                return secure_html(
                    make_page(
                        error="Something went wrong while processing "
                              "your document. Please try again."
                    ),
                    500,
                )

        return secure_html(
            make_page(
                result=result["text"],
                information=result["information"],
                preview=result["original_preview"],
                processed_preview=result["processed_preview"],
            )
        )

    except Exception:
        logger.exception("Failed to handle uploaded document")
        return secure_html(
            make_page(error="The request could not be completed."),
            500,
        )

    finally:
        await file.close()


@app.get("/health")
async def health():
    return {"status": "ok"}


@app.get("/download")
async def download_help():
    # Text downloads are generated by the browser without server-side storage.
    return PlainTextResponse(
        "Use the Download Text button after processing a document.",
        headers={"Cache-Control": "no-store"},
    )
